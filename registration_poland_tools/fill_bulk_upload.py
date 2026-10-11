#!/usr/bin/env -S uv run
"""Fill the Jamboree Poland 2027 bulk upload template from the Hitobito DB.

The mapping below follows bulk-upload-fields.md.
"""

from __future__ import annotations

import argparse
import dataclasses
import datetime
import pathlib
import re
from collections.abc import Callable
from typing import Any

import openpyxl
import pandas as pd
import wsjrdp2027
from psycopg import sql


TEMPLATE_PATH = pathlib.Path(__file__).with_name("Bulk-Upload-Template.XLSX")
SHEET_NAME = "Import Data"
HEADER_ROW = 2
FIRST_DATA_ROW = 5
# Dropdown validations of the template only reach row 205.
MAX_ROWS_PER_FILE = 201

EXPORTED_ROLES = ("YP", "UL", "IST", "CMT")

# Roles per field, colour-coded in the template header.
ALL = ["YP", "UL", "IST", "CMT"]  # red
YP_UL = ["YP", "UL"]  # blue
ADULTS = ["UL", "CMT", "IST"]  # yellow
IST = ["IST"]  # green
YP = ["YP"]  # pink

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

CONTINGENT_COUNTRY = "Germany"
CONTINGENT_NSO = "Germany - Ring deutscher Pfadfinder*innenverbände"
NATIVE_LANGUAGE = "German"
CONTINGENT_PHONE = "+4917TODO"

RDP_ADDRESS_1 = "c/o Ring deutscher Pfadfinder*innenverbände e.V. (rdp)"
RDP_ADDRESS_2 = "Chausseestraße 128/129"
RDP_POSTCODE = "10115"
RDP_CITY = "Berlin"
RDP_STATE = "Berlin"
RDP_COUNTRY = "Germany"

EMERGENCY_FIRST_NAME = "Notfalltelefon"
EMERGENCY_SURNAME = "Deutsches Kontingent"
EMERGENCY_RELATIONSHIP = "Other"
EMERGENCY_RELATIONSHIP_OTHER = "Contingent Management Team"
EMERGENCY_NATIONALITY = "Germany"
EMERGENCY_EMAIL = "notfallkontakt@worldscoutjamboree.de"

PRESELECTED_IST_TAG = "Preselected-IST"
BMT_GROUP_ID = 45

# ---------------------------------------------------------------------------
# Value lists (DB -> template)
# ---------------------------------------------------------------------------

SUB_CATEGORY = {
    "YP": "Youth Participants",
    "UL": "Adult Patrol Leaders",
    "IST": "International Service Team",
    "CMT": "Contingent Management Team",
}

GENDER = {"w": "Female", "m": "Male", "d": "Other"}

ISO_TO_COUNTRY = {
    "DE": "Germany",
    "AT": "Austria",
    "CH": "Switzerland",
    "NL": "Netherlands",
    "PL": "Poland",
    "SE": "Sweden",
    "US": "United States of America",
    "GB": "United Kingdom",
    "MY": "Malaysia",
}
ISO_TO_NATIONALITY = ISO_TO_COUNTRY | {"US": "United States"}

# Lower-case German country names and nationality adjectives -> ISO code.
GERMAN_COUNTRY_TO_ISO = {
    "deutschland": "DE",
    "österreich": "AT",
    "schweiz": "CH",
    "niederlande": "NL",
    "polen": "PL",
    "schweden": "SE",
    "usa": "US",
    "vereinigte staaten": "US",
    "großbritannien": "GB",
    "vereinigtes königreich": "GB",
    "malaysia": "MY",
}
GERMAN_NATIONALITY_TO_ISO = {
    "deutsch": "DE",
    "österreichisch": "AT",
    "österreich": "AT",
    "schweizerisch": "CH",
    "niederländisch": "NL",
    "polnisch": "PL",
    "schwedisch": "SE",
    "us-amerikanisch": "US",
    "amerikanisch": "US",
    "britisch": "GB",
    "malaysisch": "MY",
}

LANGUAGE_LEVELS = {
    "basic": "Basic",
    "conversational": "Conversational",
    "proficient": "Proficient",
    "fluent": "Fluent",
    "native": "Native",
}

SHIRT_SIZE = {
    "XXS": "XS - (Chest 80-88cm)",
    "XS": "XS - (Chest 80-88cm)",
    "S": "S (Chest 88-96cm)",
    "M": "M (96-104cm)",
    "L": "L (104-112cm)",
    "XL": "XL (112-120cm)",
    "XXL": "XXL (120-128cm)",
    "XXXL": "3XL (128-136cm)",
    "XXXXL": "3XL (128-136cm)",
}

TRAVEL_DOCUMENT_TYPE = {
    "ordinary_passport": "Ordinary passport",
    "diplomatic_passport": "Diplomatic passport",
    "service_passport": "Service passport",
    "official_passport": "Other travel document (please specify)",
    "special_passport": "Special passport",
    "other_travel_document": "Other travel document (please specify)",
}

VISA_NEEDED = {
    "no": "No, I do not need a visa",
    "yes_have_one": "Visa exists",
    "yes_need_support": "Yes, I do need a Schengen visa",
}

VISA_GENDER = {
    "m": "Male",
    "w": "Female",
    "d": "X",
    "männlich": "Male",
    "weiblich": "Female",
    "divers": "X",
}
VISA_CIVIL_STATUS = {
    "ledig": "Single",
    "verheiratet": "Married",
    "geschieden": "Divorced",
    "getrennt": "Separated",
    "verwitwet": "Widow(er)",
    "eingetragene lebenspartnerschaft": "Registered Partnership",
}
VISA_STUDENT_STATUS = {"ja": "I am a student", "nein": "I am not a student"}
VISA_FAMILY_RELATIONSHIP = {
    "ehepartner": "Spouse",
    "kind": "Child",
    "enkelkind": "Grandchild",
    "eingetragene lebenspartnerschaft": "Registered Partnership",
    "unterhaltsberechtigte vorfahren": "Dependent Ascendant",
    "sonstige": "Other",
}

DIET = {
    "omnivorous": "No special diet",
    "vegetarian": "Vegetarian",
    "vegan": "Vegan",
    "halal": "Halal",
    "kosher": "Kosher",
}

ALLERGY_SEVERITY = {
    "severe": "Severe anaphylactic reaction (could result in death)",
    "medium": "Medium response & requires medical intervention",
    "mild": "Mild response that can be self managed",
}

ALLERGIES = [
    "milk",
    "casein",
    "egg",
    "peanut",
    "tree_nuts",
    "gluten",
    "soy",
    "lupin",
    "mustard",
    "crustaceans",
    "sesame",
    "molluscs",
    "sulfite",
    "fruit",
    "allium",
    "fish",
    "seafood",
]

MEDICAL_EQUIPMENT = {
    "battery_charging_station": "I need battery charging for medical equipment or aids (e.g., Continuous positive airway (CPAP) machine user)",
    "continuous_power_supply": "I need permanent electrical supply for medical equipment or aids",
    "medication_cooling": "I need refrigeration of medication",
    "other": "Other",
}

WHEELCHAIR_TYPE = {"manual": "Manual", "electric": "Electric"}

ARRIVAL_LOCATIONS = {
    "amber_expo": "AmberExpo",
    "gdansk_ferry_terminal": "Gdansk Ferry Terminal",
    "gdansk_glowny_trainstation": "Gdansk Glowny Trainstation",
    "gdynia_ferry_terminal": "Gdynia Ferry Terminal",
    "lech_walesa_airport": "Lech Walesa Airport",
}
DEPARTURE_LOCATIONS = {
    key: value for key, value in ARRIVAL_LOCATIONS.items() if key != "amber_expo"
} | {
    "jamboree_site_pick_up_point": "Jamboree Site pick-up point",
}
TRAVEL_TYPES = {
    name.lower(): name
    for name in ("Bicycle", "Bus", "Car", "Coach", "Ferry", "Flight", "Train")
}

# Cells whose value may hold several dropdown values joined by MULTI_VALUE_SEP.
MULTI_VALUE_SEP = "; "
MULTI_VALUE_CELLS = {
    "Do you have any special needs regarding medical equipment?",
}

# ---------------------------------------------------------------------------
# Person
# ---------------------------------------------------------------------------


@dataclasses.dataclass
class Person:
    row: dict[str, Any]
    warnings: list[dict[str, Any]]

    @property
    def id(self) -> int:
        return self.row["id"]

    @property
    def role(self) -> str:
        return self.row["wsj_role"]

    def __getitem__(self, key: str) -> Any:
        return self.row.get(key)

    def jd(self, key: str) -> Any:
        value = (self.row.get("jamboree_data") or {}).get(key)
        return value.strip() or None if isinstance(value, str) else value

    def flag(self, key: str) -> bool:
        return self.jd(key) in (True, "true", "1")

    def warn(self, cell_name: str, message: str) -> None:
        self.warnings.append(
            {"id": self.id, "role": self.role, "cell": cell_name, "message": message}
        )


Value = Callable[[Person], Any] | str | None


@dataclasses.dataclass(frozen=True)
class Field:
    cell_name: str
    value: Value
    roles: list[str]


# ---------------------------------------------------------------------------
# Value functions
# ---------------------------------------------------------------------------


def yes_no(value: bool) -> str:
    return "Yes" if value else "No"


def column(key: str) -> Callable[[Person], Any]:
    return lambda p: p[key]


def jd(key: str) -> Callable[[Person], Any]:
    return lambda p: p.jd(key)


def jd_yes_no(key: str) -> Callable[[Person], str]:
    return lambda p: yes_no(p.flag(key))


def comfortable(
    key: str, yes: str = "I do feel comfortable"
) -> Callable[[Person], str]:
    return lambda p: yes if p.flag(key) else "I do not feel comfortable"


def format_date(value: Any) -> str | None:
    if not value:
        return None
    if isinstance(value, str):
        value = datetime.date.fromisoformat(value[:10])
    return value.strftime("%Y/%m/%d")


def jd_date(key: str) -> Callable[[Person], str | None]:
    return lambda p: format_date(p.jd(key))


def language(lang: str, skill: str) -> Callable[[Person], str | None]:
    # The template spells English listening as "Not Listening".
    none_label = (
        "Not Listening" if (lang, skill) == ("english", "listening") else f"Not {skill}"
    )

    def value(p: Person) -> str | None:
        level = p.jd(f"{lang}_{skill}")
        if level == "none":
            return none_label
        return LANGUAGE_LEVELS.get(level) if level else None

    return value


def allergy(key: str) -> Callable[[Person], str]:
    return jd_yes_no(f"allergy_{key}")


def allergy_severity(key: str) -> Callable[[Person], str | None]:
    def value(p: Person) -> str | None:
        if not p.flag(f"allergy_{key}"):
            return None
        return ALLERGY_SEVERITY.get(p.jd(f"allergy_{key}_severity") or "")

    return value


def translate(
    p: Person,
    cell_name: str,
    raw: str | None,
    to_iso: dict[str, str],
    iso_to_name: dict[str, str],
) -> str | None:
    if not raw:
        return None
    first = raw.split(",")[0].strip()
    for name in iso_to_name.values():
        if first.casefold() == name.casefold():
            return name
    iso = first.upper() if first.upper() in iso_to_name else to_iso.get(first.lower())
    if iso is None or iso not in iso_to_name:
        p.warn(cell_name, f"Unbekannter Wert {raw!r}")
        return None
    return iso_to_name[iso]


def country_of_residence(p: Person) -> str | None:
    return translate(p, "Country", p["country"] or "DE", {}, ISO_TO_COUNTRY)


def nationality(p: Person) -> str | None:
    raw = p.jd("nationality") or p["passport_nationality"]
    return translate(
        p, "Nationality", raw, GERMAN_NATIONALITY_TO_ISO, ISO_TO_NATIONALITY
    )


def jd_country(key: str, cell_name: str) -> Callable[[Person], str | None]:
    return lambda p: translate(
        p, cell_name, p.jd(key), GERMAN_COUNTRY_TO_ISO, ISO_TO_COUNTRY
    )


def visa_choice(
    key: str, cell_name: str, translations: dict[str, str]
) -> Callable[[Person], str | None]:
    def value(p: Person) -> str | None:
        raw = p.jd(key)
        if not raw:
            return None
        choice = translations.get(raw.casefold(), raw)
        if choice not in translations.values():
            p.warn(cell_name, f"Unbekannter Wert {raw!r}")
            return None
        return choice

    return only_with_visa(value)


def needs_visa(p: Person) -> bool:
    return p.jd("visa_needed") == "yes_need_support"


def address_1(p: Person) -> str | None:
    return " ".join(part for part in (p["street"], p["housenumber"]) if part) or None


def known_as_name(p: Person) -> str | None:
    return p.jd("known_as_name") or p["nickname"] or p["first_name"]


def gender(p: Person) -> str:
    return GENDER.get(p["gender"] or "", "Prefer not to say")


def shirt_size(p: Person) -> str | None:
    size = SHIRT_SIZE.get((p["shirt_size"] or "").strip().upper())
    if size is None:
        p.warn("T-Shirt Size WSJ", f"Keine gültige Größe: {p['shirt_size']!r}")
    return size


def preselected_ist(p: Person) -> str:
    return yes_no(
        PRESELECTED_IST_TAG in (p["tags"] or [])
        or p["primary_group_id"] == BMT_GROUP_ID
    )


def guardian_given_name(p: Person) -> str | None:
    name = (p["additional_contact_name_a"] or "").strip()
    return name.rsplit(" ", 1)[0] if " " in name else None


def guardian_surname(p: Person) -> str | None:
    name = (p["additional_contact_name_a"] or "").strip()
    return name.rsplit(" ", 1)[-1] if name else None


def travel_doc_other(p: Person) -> str | None:
    if p.jd("travel_document_type") == "other_travel_document":
        p.warn("Travel Doc Other", "Art des anderen Reisedokuments von Hand nachtragen")
    return (
        "German ID card (Personalausweis)"
        if p.jd("travel_document_type") == "official_passport"
        else None
    )


def free_texts(*keys: str) -> Callable[[Person], str | None]:
    return lambda p: MULTI_VALUE_SEP.join(t for k in keys if (t := p.jd(k))) or None


def has_dietary_requirements(p: Person) -> str:
    return yes_no(
        p["diet"] not in (None, "omnivorous")
        or bool(p.jd("dietary_requirements") or p.jd("dietary_details"))
    )


def special_diet(p: Person) -> str | None:
    diet = p["diet"] or "omnivorous"
    if diet == "omnivorous" and (
        p.jd("dietary_requirements") or p.jd("dietary_details")
    ):
        return "Other"
    return DIET.get(diet)


def has_allergies(p: Person) -> str:
    return yes_no(any(p.flag(f"allergy_{a}") for a in ALLERGIES))


def has_text(key: str) -> Callable[[Person], str]:
    return lambda p: yes_no(bool((p[key] or "").strip()))


def column_yes_no(key: str) -> Callable[[Person], str]:
    return lambda p: yes_no(bool(p[key]))


def medical_equipment(p: Person) -> str | None:
    needs = p.jd("medical_equipment_needs") or []
    return MULTI_VALUE_SEP.join(MEDICAL_EQUIPMENT.get(n, n) for n in needs) or None


def wheelchair_type(p: Person) -> str | None:
    return (
        WHEELCHAIR_TYPE.get(p.jd("wheelchair_type") or "")
        if p.flag("uses_wheelchair")
        else None
    )


def only_with_visa(value: Callable[[Person], Any]) -> Callable[[Person], Any]:
    return lambda p: value(p) if needs_visa(p) else None


def visa_fingerprints(p: Person) -> str | None:
    if not needs_visa(p) or not p.flag("visa_fingerprints_collected"):
        return None
    return (
        MULTI_VALUE_SEP.join(
            part
            for part in (
                format_date(p.jd("visa_fingerprints_date")),
                p.jd("visa_fingerprints_visa_number"),
            )
            if part
        )
        or None
    )


def visa_family_member(p: Person) -> str | None:
    if not needs_visa(p):
        return None
    return "Yes" if p.jd("visa_family_eu_citizen_relationship") else None


# ---------------------------------------------------------------------------
# Mapping: one entry per template column, in template order
# ---------------------------------------------------------------------------

FIELDS: list[Field] = [
    # General
    Field("Sub Category", lambda p: SUB_CATEGORY[p.role], ALL),
    Field("Please agree to the GDPR agreement", jd_yes_no("gdpr_agreement"), ALL),
    Field("Contingent Country Name", CONTINGENT_COUNTRY, ALL),
    Field("Contingent NSO Name", CONTINGENT_NSO, ALL),
    Field("Patrol ID", column("patrol"), YP_UL),
    # Person
    Field("First Name", column("first_name"), ALL),
    Field("Last Name", column("last_name"), ALL),
    Field("Surname at birth (former family name(s))", None, ALL),
    Field("Known as name", known_as_name, ALL),
    Field("Mother's First Name", jd("mothers_first_name"), ADULTS),
    Field("Father's First Name", jd("fathers_first_name"), ADULTS),
    Field("Mother's surname at birth", jd("mothers_surname_at_birth"), ADULTS),
    Field(
        "Date of Birth (dd.mm.yyy) [dd/MM/yyyy]",
        lambda p: format_date(p["birthday"]),
        ALL,
    ),
    Field("Place of Birth", jd("place_of_birth"), ADULTS),
    Field("Country of birth", jd_country("country_of_birth", "Country of birth"), ALL),
    Field("Gender (please select)", gender, ALL),
    Field("Religion(s)", None, ALL),
    Field("Nationality", nationality, ALL),
    Field("Please state your National Identity Number (if applicable)", None, ALL),
    # Address & contact
    Field("Address 1", address_1, ALL),
    Field("Address 2", None, ALL),
    Field("City", column("town"), ALL),
    Field("State County", None, ALL),
    Field("Postcode", column("zip_code"), ALL),
    Field("Country", country_of_residence, ALL),
    Field("Email Address / Adresse e-mail", column("email"), ALL),
    Field("Phone number (incl. country code +XX)", CONTINGENT_PHONE, ADULTS),
    Field(
        "Do you need an Adult that will accompany you in the WSJ as a carer (AC)? "
        "Please fill out ONLY IF an approved individual support arrangement is required.",
        "No",
        ALL,
    ),
    Field("AC - First name(s) (Given name(s)) as in the travel document", None, ALL),
    Field("AC - Surname (family name) - as in the travel document", None, ALL),
    Field("AC - Email address", None, ALL),
    # Languages
    Field("What is your native language?", NATIVE_LANGUAGE, ALL),
    Field(
        "What is your English reading language proficiency level",
        language("english", "reading"),
        ALL,
    ),
    Field(
        "What is your English writing language proficiency level",
        language("english", "writing"),
        ALL,
    ),
    Field(
        "What is your English speaking language proficiency level",
        language("english", "speaking"),
        ALL,
    ),
    Field(
        "What is your English listening language proficiency level",
        language("english", "listening"),
        ALL,
    ),
    Field(
        "What is your French reading language proficiency level",
        language("french", "reading"),
        ADULTS,
    ),
    Field(
        "What is your French writing language proficiency level",
        language("french", "writing"),
        ADULTS,
    ),
    Field(
        "What is your French speaking language proficiency level",
        language("french", "speaking"),
        ADULTS,
    ),
    Field(
        "What is your French listening language proficiency level",
        language("french", "listening"),
        ADULTS,
    ),
    Field(
        "What is your Spanish reading language proficiency level",
        language("spanish", "reading"),
        ADULTS,
    ),
    Field(
        "What is your Spanish writing language proficiency level",
        language("spanish", "writing"),
        ADULTS,
    ),
    Field(
        "What is your Spanish speaking language proficiency level",
        language("spanish", "speaking"),
        ADULTS,
    ),
    Field(
        "What is your Spanish listening language proficiency level",
        language("spanish", "listening"),
        ADULTS,
    ),
    Field(
        "What is your Arabic reading language proficiency level",
        language("arabic", "reading"),
        ADULTS,
    ),
    Field(
        "What is your Arabic writing language proficiency level",
        language("arabic", "writing"),
        ADULTS,
    ),
    Field(
        "What is your Arabic speaking language proficiency level",
        language("arabic", "speaking"),
        ADULTS,
    ),
    Field(
        "What is your Arabic listening language proficiency level",
        language("arabic", "listening"),
        ADULTS,
    ),
    # Comfort
    Field(
        "Do you feel comfortable being on Water? (Sailing/Kayaking)",
        comfortable("comfortable_on_water"),
        ALL,
    ),
    Field(
        "Do you feel comfortable being in Water? (Swimming)",
        comfortable("comfortable_in_water"),
        ALL,
    ),
    Field(
        "Do you feel comfortable without balance?",
        comfortable("comfortable_without_balance"),
        ALL,
    ),
    Field(
        "Do you feel comfortable in a large crowd?",
        comfortable("comfortable_in_crowd", yes="I feel comfortable"),
        ALL,
    ),
    Field(
        "Do you feel comfortable with flashing lights?",
        comfortable("comfortable_flashing_lights"),
        ALL,
    ),
    Field(
        "Do you feel comfortable with strong noise?",
        comfortable("comfortable_loud_noise"),
        ALL,
    ),
    Field(
        "Do you feel comfortable in a confined space?",
        comfortable("comfortable_confined_space"),
        ALL,
    ),
    # Other
    Field("Prior WSJ Experience", jd("prior_wsj_experience"), ADULTS),
    Field("T-Shirt Size WSJ", shirt_size, ADULTS),
    Field("Pre-selected IST for Contingent Pavilions", preselected_ist, IST),
    Field("Pre-selected IST for Food Houses", preselected_ist, IST),
    Field(
        "If you arrive by car or have a car available during your stay, "
        "please state your registration number (number plate) here.",
        jd("vehicle_registration_number"),
        ADULTS,
    ),
    # Parental authority / legal guardian (contact A, address of the rdp office)
    Field("Parental authority / Legal guardian - Given name", guardian_given_name, YP),
    Field(
        "Parental authority / Legal guardian - Surname (as in legal document)",
        guardian_surname,
        YP,
    ),
    Field("Parental authority / Legal guardian - Relationship", "Parent", YP),
    Field("Parent/Guardian Other", None, YP),
    Field("Parental Authority / Legal Guardian -Nationality", None, YP),
    Field(
        "Parental Authority / Legal Guardian - Email address",
        column("additional_contact_email_a"),
        YP,
    ),
    Field("Parental Authority / Legal Guardian - Phone number", CONTINGENT_PHONE, YP),
    Field("Parental Authority / Legal Guardian - Language", None, YP),
    Field("Parental Authority / Legal Guardian - Home address 1", RDP_ADDRESS_1, YP),
    Field("Parental Authority / Legal Guardian - Home address 2", RDP_ADDRESS_2, YP),
    Field("Parental Authority / Legal Guardian - City", RDP_CITY, YP),
    Field("Parental Authority / Legal Guardian - State", RDP_STATE, YP),
    Field("Parental Authority / Legal Guardian - Postcode", RDP_POSTCODE, YP),
    Field("Parental Authority / Legal Guardian - Country", RDP_COUNTRY, YP),
    # Travel document
    Field(
        "Travel Document Type",
        lambda p: TRAVEL_DOCUMENT_TYPE.get(p.jd("travel_document_type") or ""),
        ALL,
    ),
    Field("Travel Doc Other", travel_doc_other, ALL),
    Field("Travel Document Number", jd("travel_document_number"), ALL),
    Field("Travel Document Date of Issue", jd_date("travel_document_issue_date"), ALL),
    Field(
        "Travel Document Expiration Date",
        jd_date("travel_document_expiration_date"),
        ALL,
    ),
    Field(
        "Travel Document Issued By (Country)",
        jd_country("travel_document_country", "Travel Document Issued By (Country)"),
        ALL,
    ),
    # Visa
    Field(
        "Do you need a Visa to enter Poland?",
        lambda p: VISA_NEEDED.get(p.jd("visa_needed") or "no"),
        ALL,
    ),
    Field("Visa Information", jd("visa_information"), ALL),
    Field("Place of Birth", only_with_visa(jd("place_of_birth")), ALL),
    Field(
        "Country of Birth",
        only_with_visa(jd_country("country_of_birth", "Country of Birth")),
        ALL,
    ),
    Field(
        "Nationality at birth if different than the country of birth",
        only_with_visa(jd("visa_nationality_at_birth")),
        ALL,
    ),
    Field(
        "Other / Further nationalities",
        only_with_visa(jd("visa_other_nationalities")),
        ALL,
    ),
    Field(
        "Gender as in the travel document",
        visa_choice(
            "visa_gender_as_in_travel_document",
            "Gender as in the travel document",
            VISA_GENDER,
        ),
        ALL,
    ),
    Field(
        "Civil Status",
        visa_choice("visa_civil_status", "Civil Status", VISA_CIVIL_STATUS),
        ALL,
    ),
    Field("Civil Status Other", only_with_visa(jd("visa_civil_status_other")), ALL),
    Field(
        "Current student status",
        visa_choice(
            "visa_current_student_status", "Current student status", VISA_STUDENT_STATUS
        ),
        ALL,
    ),
    Field(
        "Do you have a family member who is a citizen of CH or a EU/EEA country?",
        visa_family_member,
        ALL,
    ),
    Field(
        "Employer’s address. For students, please state the name and address of your educational establishment",
        only_with_visa(jd("visa_employer_name")),
        ALL,
    ),
    Field(
        "Employer's / Educational institution's address 1",
        only_with_visa(jd("visa_employer_address_1")),
        ALL,
    ),
    Field(
        "Employer's / Educational institution's address 2",
        only_with_visa(jd("visa_employer_address_2")),
        ALL,
    ),
    Field(
        "Employer's / Educational institution's City",
        only_with_visa(jd("visa_employer_city")),
        ALL,
    ),
    Field(
        "Employer's / Educational institution's Postcode",
        only_with_visa(jd("visa_employer_postcode")),
        ALL,
    ),
    Field(
        "Employer's / Educational institution's State",
        only_with_visa(jd("visa_employer_state")),
        ALL,
    ),
    Field(
        "Employer's / Educational institution's Country",
        only_with_visa(
            jd_country(
                "visa_employer_country",
                "Employer's / Educational institution's Country",
            )
        ),
        ALL,
    ),
    Field(
        "Please state your family relationship with a EU, EEA or CH citizen (FAM EU II)",
        visa_choice(
            "visa_family_eu_citizen_relationship",
            "Please state your family relationship with a EU, EEA or CH citizen (FAM EU II)",
            VISA_FAMILY_RELATIONSHIP,
        ),
        ALL,
    ),
    Field(
        "Please specify your family relation to an EU, EEA or CH citizen (FAM EU III)",
        only_with_visa(jd("visa_family_eu_citizen_relation")),
        ALL,
    ),
    Field(
        "Residence in a country other than the country of current nationality",
        only_with_visa(jd_yes_no("visa_residence_other_than_nationality")),
        ALL,
    ),
    Field(
        "If residence other than current nationality",
        only_with_visa(
            jd_country(
                "visa_residence_country", "If residence other than current nationality"
            )
        ),
        ALL,
    ),
    Field(
        "Schengen Member State of first entry",
        only_with_visa(
            jd_country(
                "visa_schengen_first_entry", "Schengen Member State of first entry"
            )
        ),
        ALL,
    ),
    Field(
        "Fingerprints collected previously for the purpose of applying for a Schengen visa",
        only_with_visa(jd_yes_no("visa_fingerprints_collected")),
        ALL,
    ),
    Field(
        "If yes, please state the date when this took place and the number of your visa sticker "
        "(if known) (Fingerprnts II)",
        visa_fingerprints,
        ALL,
    ),
    # Arrival / departure
    Field("Arrival and Departure Information", None, ALL),
    Field("Arrival", None, ALL),
    Field(
        "Arrival Location",
        lambda p: ARRIVAL_LOCATIONS.get(p.jd("arrival_location")),
        ALL,
    ),
    Field("Arrival Date", jd_date("arrival_date"), ALL),
    Field("Arrival time", jd("arrival_time"), ALL),
    Field("Arrival Type", lambda p: TRAVEL_TYPES.get(p.jd("arrival_type")), ALL),
    Field("Arrival Number", jd("arrival_number"), ALL),
    Field("Departure", None, ALL),
    Field(
        "Departure Location",
        lambda p: DEPARTURE_LOCATIONS.get(p.jd("departure_location")),
        ALL,
    ),
    Field("Departure Date", jd_date("departure_date"), ALL),
    Field("Departure Time", jd("departure_time"), ALL),
    Field("Departure Type", lambda p: TRAVEL_TYPES.get(p.jd("departure_type")), ALL),
    Field("Departure Number", jd("departure_number"), ALL),
    # Emergency contact: fixed CMT contact
    Field("Emergency Contact", None, ALL),
    Field("Emergency Contact - First (Given) Name", EMERGENCY_FIRST_NAME, ALL),
    Field("Emergency Contact - Surname", EMERGENCY_SURNAME, ALL),
    Field("Emergency Contact - Relationship", EMERGENCY_RELATIONSHIP, ALL),
    Field("Emergency Contact Relationship Other", EMERGENCY_RELATIONSHIP_OTHER, ALL),
    Field("Emergency Contact - Nationality", EMERGENCY_NATIONALITY, ALL),
    Field("Emergency Contact Email Address", EMERGENCY_EMAIL, ALL),
    Field("Emergency Contact - Phone Number", CONTINGENT_PHONE, ALL),
    Field("Emergency Contact - Address 1", RDP_ADDRESS_1, ALL),
    Field("Emergency Contact - Address 2", RDP_ADDRESS_2, ALL),
    Field("Emergency Contact - City", RDP_CITY, ALL),
    Field("Emergency Contact - State", RDP_STATE, ALL),
    Field("Emergency Contact - Postcode", RDP_POSTCODE, ALL),
    Field("Emergency Contact - Country", RDP_COUNTRY, ALL),
    # Diet
    Field("Health & Diet", None, ALL),
    Field(
        "Do you have any dietary requirements or preferences?",
        has_dietary_requirements,
        ALL,
    ),
    Field(
        "Please select your dietary requirements or preferences (special diet II)",
        special_diet,
        ALL,
    ),
    Field(
        "Please specify your special diet (special diet III)",
        free_texts("dietary_requirements", "dietary_details"),
        ALL,
    ),
    # Allergies
    Field("Do you have any allergies?", has_allergies, ALL),
    Field("Do you have an allergy to Milk (Cows Milk)?", allergy("milk"), ALL),
    Field(
        "What is your Allergy Severity to Milk (Cows Milk)?",
        allergy_severity("milk"),
        ALL,
    ),
    Field("Do you have an allergy to Casein?", allergy("casein"), ALL),
    Field("What is your Allergy Severity to Casein?", allergy_severity("casein"), ALL),
    Field("Do you have an allergy to Egg?", allergy("egg"), ALL),
    Field("What is your Allergy severity to Egg?", allergy_severity("egg"), ALL),
    Field("Do you have an allergy to Peanuts?", allergy("peanut"), ALL),
    Field("What is your Allergy severity to Peanuts?", allergy_severity("peanut"), ALL),
    Field(
        "Do you have an allergy to Tree Nuts (almonds, hazelnuts, walnuts, Brazil nuts, cashews, "
        "pecans, pistachios, macadamia nuts)?",
        allergy("tree_nuts"),
        ALL,
    ),
    Field(
        "What is your allergy severity to Tree Nuts (almonds, hazelnuts, walnuts, Brazil nuts, "
        "cashews, pecans, pistachios, macadamia nuts)?",
        allergy_severity("tree_nuts"),
        ALL,
    ),
    Field(
        "Do you have an allergy to Cereals containing gluten(wheat, rye, barley, oats)?",
        allergy("gluten"),
        ALL,
    ),
    Field(
        "What is your Allergy severity to Cereals containing gluten (Wheat, Rye, Barley, Oats)?",
        allergy_severity("gluten"),
        ALL,
    ),
    Field("Do you have an allergy to Soybeans?", allergy("soy"), ALL),
    Field("What is your Allergy Severity to Soybeans?", allergy_severity("soy"), ALL),
    Field("Do you have an allergy to Lupin?", allergy("lupin"), ALL),
    Field("What is your Allergy Severity to Lupin?", allergy_severity("lupin"), ALL),
    Field("Do you have an allergy to Mustard?", allergy("mustard"), ALL),
    Field(
        "What is your Allergy Severity to Mustard?", allergy_severity("mustard"), ALL
    ),
    Field(
        "Do you have an allergy to Crustaceans (Prawns, Crabs, Lobsters)?",
        allergy("crustaceans"),
        ALL,
    ),
    Field(
        "What is your Allergy Severity to Crustaceans (Prawns, Crabs, Lobsters)?",
        allergy_severity("crustaceans"),
        ALL,
    ),
    Field("Do you have an allergy to Sesame?", allergy("sesame"), ALL),
    Field("What is your allergy severity to Sesame?", allergy_severity("sesame"), ALL),
    Field(
        "Do you have an allergy to Molluscs (Mussels, Oysters)?",
        allergy("molluscs"),
        ALL,
    ),
    Field(
        "What is your Allergy Severity to Molluscs (Mussels, Oysters)?",
        allergy_severity("molluscs"),
        ALL,
    ),
    Field(
        "Do you have an allergy to Sulphur Dioxide & Sulphites (if the Sulphur Dioxide and Sulphites "
        "are at a concentration of more than ten parts per million)?",
        allergy("sulfite"),
        ALL,
    ),
    Field(
        "What is your Allergy Severity to Sulphur Dioxide & Sulphites (If the Sulphur Dioxide and "
        "Sulphites are at a concentration of more than ten parts per million)?",
        allergy_severity("sulfite"),
        ALL,
    ),
    Field("Do you have an allergy to Fruit?", allergy("fruit"), ALL),
    Field("What is your Allergy Severity to Fruit?", allergy_severity("fruit"), ALL),
    Field(
        "Do you have an allergy to Alliums (Onion, Garlic, Chive, ETC)?",
        allergy("allium"),
        ALL,
    ),
    Field(
        "What is your Allergy Severity to Alliums (Onions, Garlic, Chive, ETC)?",
        allergy_severity("allium"),
        ALL,
    ),
    Field("Do you have an allergy to Fish?", allergy("fish"), ALL),
    Field("What is your allergy severity to Fish?", allergy_severity("fish"), ALL),
    Field("Do you have an allergy to Shellfish?", allergy("seafood"), ALL),
    Field(
        "What is your Allergy Severity to Shellfish?", allergy_severity("seafood"), ALL
    ),
    # Medical
    Field(
        "Do you suffer from any chronic diseases?",
        has_text("medical_preexisting_conditions"),
        ALL,
    ),
    Field(
        "If you suffer from chronic diseas(es) please specify",
        column("medical_preexisting_conditions"),
        ALL,
    ),
    Field(
        "Please specify all medicines you consume regularly, along with dosage "
        "(trade names of medications along with chemical names)",
        column("medical_continuous_medication"),
        ALL,
    ),
    Field(
        "Does your medication need to be refrigerated?",
        jd_yes_no("refrigerated_medication_needed"),
        ALL,
    ),
    Field(
        "Do you have any religious or clerical beliefs that influence health decisions?",
        jd("religious_health_beliefs"),
        ALL,
    ),
    Field(
        "Are you vaccinated with Tetanus?",
        column_yes_no("medical_stiko_vaccinations"),
        ALL,
    ),
    Field("Tetanus Vaccination when", None, ALL),
    Field(
        "Are you vaccinated against Diphteria?",
        column_yes_no("medical_stiko_vaccinations"),
        ALL,
    ),
    Field("Diphteria Vaccination When", None, ALL),
    Field("What is your blood type?", None, ALL),
    Field(
        "Do you have any special needs regarding medical equipment?",
        medical_equipment,
        ALL,
    ),
    Field(
        "Please specify your needs re your medical equipment",
        jd("medical_equipment_details"),
        ALL,
    ),
    # Mobility
    Field(
        "Do you have any mobility restrictions?", jd_yes_no("mobility_restricted"), ALL
    ),
    Field("Are you using a Wheelchair?", jd_yes_no("uses_wheelchair"), ALL),
    Field("Is the Wheelchair Manual or Electric?", wheelchair_type, ALL),
    Field(
        "Are you using a Cane/Stick/Frame for mobility?",
        jd_yes_no("uses_mobility_frame"),
        ALL,
    ),
    Field(
        "Do you have any other mobility needs?", jd_yes_no("other_mobility_needs"), ALL
    ),
    Field("Further Mobility Needs", jd("other_mobility_details"), ALL),
    Field("Is a larger tent required?", None, ALL),
    # Non-visible disabilities
    Field(
        "Do you have any non-visible disability?",
        jd_yes_no("non_visible_disability"),
        ALL,
    ),
    Field("Are you blind or have low-vision?", None, ALL),
    Field("Do you have any other significant visual impairment?", None, ALL),
    Field("Are you hard of hearing or deaf?", None, ALL),
    Field(
        "Do you struggle with any of the following: loud noises, bright lights, crowds?",
        jd_yes_no("sensory_sensitivities"),
        ALL,
    ),
    Field(
        "Do you have any other non-visible disability that could impact participation or require additional support?",
        jd_yes_no("non_visible_disability"),
        ALL,
    ),
    Field(
        "Please specify the condition/disability and the medication/adaptations needed",
        jd("non_visible_disability_details"),
        ALL,
    ),
]

# ---------------------------------------------------------------------------
# Database
# ---------------------------------------------------------------------------

PEOPLE_QUERY = """
SELECT
  p.id, p.wsj_role, p.status, p.primary_group_id, p.patrol,
  p.first_name, p.last_name, p.nickname, p.email, p.birthday, p.gender,
  p.street, p.housenumber, p.zip_code, p.town, p.country,
  p.shirt_size, p.diet, p.passport_nationality,
  p.additional_contact_name_a, p.additional_contact_email_a,
  p.medical_preexisting_conditions, p.medical_continuous_medication,
  p.medical_stiko_vaccinations,
  p.jamboree_data,
  ARRAY(
    SELECT t.name FROM taggings g JOIN tags t ON t.id = g.tag_id
    WHERE g.taggable_type = 'Person' AND g.taggable_id = p.id
  ) AS tags
FROM people p
WHERE p.status = 'confirmed' AND p.wsj_role = ANY({roles}) {id_filter}
ORDER BY array_position({roles}, p.wsj_role), p.id
{limit}
"""


def load_people(
    conn: Any, *, ids: list[int] | None, limit: int | None
) -> list[dict[str, Any]]:
    query = sql.SQL(PEOPLE_QUERY).format(
        roles=sql.Literal(list(EXPORTED_ROLES)),
        id_filter=sql.SQL("AND p.id = ANY({})").format(sql.Literal(ids))
        if ids
        else sql.SQL(""),
        limit=sql.SQL("LIMIT {}").format(sql.Literal(limit)) if limit else sql.SQL(""),
    )
    return wsjrdp2027.pg_select_dict_rows(conn, query)


def count_confirmed_without_role(conn: Any) -> int:
    rows = wsjrdp2027.pg_select_dict_rows(
        conn,
        "SELECT count(*) AS n FROM people WHERE status = 'confirmed' AND wsj_role IS NULL",
    )
    return rows[0]["n"]


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------


def normalize(text: str) -> str:
    return " ".join(str(text).split())


def map_fields_to_columns(ws: Any) -> list[tuple[Field, int]]:
    """Columns per field, matched by header text (duplicate headers in template order)."""
    header_columns: dict[str, list[int]] = {}
    for cell in ws[HEADER_ROW]:
        if cell.value is not None:
            header_columns.setdefault(normalize(cell.value), []).append(cell.column)

    result = []
    missing = []
    for field in FIELDS:
        columns = header_columns.get(normalize(field.cell_name))
        if not columns:
            missing.append(field.cell_name)
            continue
        result.append((field, columns.pop(0)))
    unmapped = [name for name, cols in header_columns.items() for _ in cols]
    if missing or unmapped:
        raise RuntimeError(
            f"Mapping does not match template: missing={missing} unmapped={unmapped}"
        )
    return result


def load_dropdown_values(wb: Any, ws: Any) -> dict[int, set[str]]:
    """Allowed values per column, read from the template's data validations."""
    allowed: dict[int, set[str]] = {}
    for dv in ws.data_validations.dataValidation:
        match = re.fullmatch(
            r"'?(.+?)'?!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", dv.formula1 or ""
        )
        if not match:
            continue
        sheet, col, first, _, last = match.groups()
        values = {
            str(c.value)
            for (c,) in wb[sheet][f"{col}{first}:{col}{last}"]
            if c.value is not None
        }
        for cell_range in dv.sqref.ranges:
            for column in range(cell_range.min_col, cell_range.max_col + 1):
                allowed[column] = values
    return allowed


def compute_value(field: Field, person: Person) -> Any:
    if person.role not in field.roles:
        return None
    value = field.value(person) if callable(field.value) else field.value
    if isinstance(value, str):
        value = value.strip() or None
    return value


def is_allowed(field: Field, value: Any, allowed: set[str]) -> bool:
    if normalize(field.cell_name) in {normalize(n) for n in MULTI_VALUE_CELLS}:
        return all(part in allowed for part in str(value).split(MULTI_VALUE_SEP))
    return str(value) in allowed


def write_file(
    ctx: wsjrdp2027.WsjRdpContext, people: list[Person], part: int, n_parts: int
) -> pathlib.Path:
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb[SHEET_NAME]
    columns = map_fields_to_columns(ws)
    dropdowns = load_dropdown_values(wb, ws)

    for row_index, person in enumerate(people, start=FIRST_DATA_ROW):
        for field, column in columns:
            value = compute_value(field, person)
            if (
                value is not None
                and column in dropdowns
                and not is_allowed(field, value, dropdowns[column])
            ):
                person.warn(field.cell_name, f"Wert nicht in Dropdown-Liste: {value!r}")
                value = None
            ws.cell(row=row_index, column=column, value=value)

    suffix = f"__part{part:02d}of{n_parts:02d}" if n_parts > 1 else ""
    out_path = ctx.make_out_path(f"bulk_upload__{{{{ filename_suffix }}}}{suffix}.xlsx")
    wb.save(out_path)
    ctx.logger.info("Wrote %s (%d people)", out_path, len(people))
    return out_path


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument(
        "--id",
        dest="ids",
        type=int,
        action="append",
        help="Only export this person id (repeatable)",
    )
    p.add_argument("--limit", type=int, help="Export at most this many people")
    ctx = wsjrdp2027.WsjRdpContext(__file__=__file__, argument_parser=p)
    args = ctx.parsed_args

    with ctx:
        ctx.configure_log_file(
            ctx.make_out_path("bulk_upload__{{ filename_suffix }}.log")
        )
        with ctx.hitobito_psycopg_connection(read_only=True) as conn:
            rows = load_people(conn, ids=args.ids, limit=args.limit)
            if n := count_confirmed_without_role(conn):
                ctx.logger.warning(
                    "%d bestätigte Personen ohne wsj_role werden nicht exportiert", n
                )

        warnings: list[dict[str, Any]] = []
        people = [Person(row, warnings) for row in rows]
        ctx.logger.info("Loaded %d people", len(people))

        chunks = [
            people[i : i + MAX_ROWS_PER_FILE]
            for i in range(0, len(people), MAX_ROWS_PER_FILE)
        ]
        for part, chunk in enumerate(chunks, start=1):
            write_file(ctx, chunk, part, len(chunks))

        if warnings:
            warnings_path = ctx.make_out_path(
                "bulk_upload__{{ filename_suffix }}__warnings.xlsx"
            )
            wsjrdp2027.write_dataframe_to_xlsx(pd.DataFrame(warnings), warnings_path)
            ctx.logger.warning("%d Warnungen, siehe %s", len(warnings), warnings_path)


if __name__ == "__main__":
    main()
