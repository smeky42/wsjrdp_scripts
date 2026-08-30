#!/usr/bin/env -S uv run
"""Fake DATEV Buchungsstapel (DTVF) files from Primanota Excel exports.

SUPERSEDED: a REAL 2025 DTVF export exists now
(External_Data/DATEV_FY2025_Export_*/), and the
re-import runbook uses it. This script is kept for its history and as a fixture
generator; its output is archived under External_Data/Fake_DATEV_Export_2025/.
Do not feed the generated files into the database alongside the real export --
the same bookings would arrive twice under different GUIDs.

Purpose (historical): development fixtures for the Buchungsstapel-based
importer, back when the accountant's DTVF export covered fiscal year 2026 only
and 2025 existed as Primanota Excel reports only. This script converts each
Primanota .xlsx into a DTVF CSV that structurally mimics a real DATEV export
(see datev_import.md for the format comparison):

  * one file per Primanota/Stapel, named DTVF_Buchungsstapel_<stamp>_NNNNN.csv,
    NNNNN ordered by (fiscal year, period, primanota number);
  * metadata header line 1 built via wsjrdp2027.datev.build_buchungsstapel_header_line
    with Kennzeichen "DTVF", Bezeichnung = the Primanota label, Datum von/bis =
    the calendar month of the Primanota period, and the observed-but-undocumented
    header fields 24 ("MP") / 26 (per-batch number) mimicked deterministically;
  * CP1252 encoding, CRLF line endings, amounts with two decimals -- like the
    real export (the Excel mojibake of the 2025 Einzug texts is preserved
    verbatim: that IS what DATEV stores);
  * "Buchungs GUID" (field 103) = **deterministic fake** UUIDv5 over
    (Beraternummer, Mandantennummer, Primanota-Nummer, Nr.), stable across
    runs -- mirroring the assumption that real DATEV GUIDs are permanent unique
    keys. Real GUIDs cannot be reconstructed from the Excel;
  * Beleglink / Beleginfo / Zusatzinformation stay EMPTY (not reconstructible;
    the Excel only carries the bare "ZI" marker).

What the Excel cannot provide and how it is faked (decided 2026-08-21):
  GUID -> deterministic UUIDv5; Beleg columns -> empty; header "erzeugt am" ->
  last day of the period, 12:00 (deterministic; the field must not be relied on
  anyway); header Herkunft -> majority HK of the sheet; field 26 -> 6-digit
  hash of the Primanota number.

Usage:
  ./accounting_tools/fake_datev_buchungsstapel.py                 # 2025 sheets
  ./accounting_tools/fake_datev_buchungsstapel.py \
      --glob 'External_Data/DATEV_Primanota_Examples/*.xlsx'
"""

from __future__ import annotations

import argparse as _argparse
import calendar as _calendar
import datetime as _datetime
import hashlib as _hashlib
import logging as _logging
import pathlib as _pathlib
import re as _re
import uuid as _uuid
from decimal import Decimal as _Decimal

import openpyxl as _openpyxl
from wsjrdp2027 import datev as _datev


_LOGGER = _logging.getLogger("fake_datev_buchungsstapel")

# Title-row parsing -- same conventions as accounting_tools/import_datev_primanota.py.
_RE_CONSULTANT_CLIENT_YEAR = _re.compile(r"(\d+)/(\d+)/(\d{4})")
_RE_PRIMANOTA = _re.compile(r"Primanota\s+(\S+)(?:\s+(.*))?$")
_RE_PERIOD = _re.compile(r"(\d{2})-(\d{4})")
_ABSTIMMSUMME_TEXT = "Abstimmsumme"

# Deterministic namespace for the fake Buchungs GUIDs.
_GUID_NAMESPACE = _uuid.uuid5(
    _uuid.NAMESPACE_URL, "https://wsjrdp.de/fake-datev-buchungsstapel"
)


def _cp1252_safe(value: str) -> str:
    """Best-effort CP1252 transliteration (the DTVF files are CP1252)."""
    try:
        value.encode("cp1252")
        return value
    except UnicodeEncodeError:
        return value.encode("cp1252", errors="replace").decode("cp1252")


def _cell_str(value) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    # openpyxl yields ints as e.g. 9500 but floats as 9500.0 for numeric cells
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def _parse_date(value) -> _datetime.date:
    if isinstance(value, _datetime.datetime):
        return value.date()
    if isinstance(value, _datetime.date):
        return value
    return _datetime.datetime.strptime(str(value).strip(), "%d.%m.%Y").date()


def _amount_de(value) -> str:
    """DATEV Umsatz formatting: sign-less, comma decimal, always 2 decimals."""
    amount = _Decimal(str(value))
    if amount < 0:
        _LOGGER.warning("Negative Umsatz %s -- writing the absolute value", value)
        amount = -amount
    return f"{amount:.2f}".replace(".", ",")


def _read_primanota(path: _pathlib.Path) -> dict:
    wb = _openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    title = str(rows[0][0] or "")
    m = _RE_CONSULTANT_CLIENT_YEAR.search(title)
    if not m:
        raise ValueError(f"{path.name}: Berater/Mandant/Jahr nicht im Titel: {title!r}")
    consultant, client, fiscal_year = m.group(1), m.group(2), int(m.group(3))
    pm = _RE_PRIMANOTA.search(title)
    if not pm:
        raise ValueError(f"{path.name}: Primanota-Nummer nicht im Titel: {title!r}")
    primanota_number = pm.group(1)
    primanota_label = (pm.group(2) or "").strip()
    per = _RE_PERIOD.search(primanota_number)
    if not per:
        raise ValueError(f"{path.name}: Periode nicht in {primanota_number!r}")
    period = _datetime.date(int(per.group(2)), int(per.group(1)), 1)

    header = [str(c) if c is not None else "" for c in rows[1]]
    idx = {name: i for i, name in enumerate(header)}

    def cell(row, name):
        i = idx.get(name)
        return row[i] if i is not None and i < len(row) else None

    data = []
    for row in rows[2:]:
        nr = cell(row, "Nr.")
        text = _cell_str(cell(row, "Buchungstext"))
        if nr in (None, "") or text == _ABSTIMMSUMME_TEXT:
            continue  # Abstimmsumme / Gruppensumme / blank rows carry no Nr.
        data.append(
            {
                "nr": int(nr),
                "umsatz": cell(row, "Umsatz"),
                "sh": _cell_str(cell(row, "S/H")) or "S",
                "wkz": _cell_str(cell(row, "WKZ")),
                "kurs": _cell_str(cell(row, "Kurs")),
                "umsatz_bw": _cell_str(cell(row, "Umsatz-BW")),
                "wkz_bw": _cell_str(cell(row, "WKZ-BW")),
                "konto": _cell_str(cell(row, "Konto")),
                "gegenkonto": _cell_str(cell(row, "Gegenkonto")),
                "bu": _cell_str(cell(row, "BU")),
                "gu": _cell_str(cell(row, "GU")),
                "belegfeld1": _cell_str(cell(row, "Belegfeld 1")),
                "belegfeld2": _cell_str(cell(row, "Belegfeld 2")),
                "datum": _parse_date(cell(row, "Datum")),
                "leistungsdatum": cell(row, "Leistungsdatum"),
                "kost1": _cell_str(cell(row, "KOST1")),
                "kost2": _cell_str(cell(row, "KOST2")),
                "kost_menge": _cell_str(cell(row, "KOST-Menge")),
                "skonto": _cell_str(cell(row, "Skonto")),
                "text": _cell_str(cell(row, "Buchungstext")),
                "hk": _cell_str(cell(row, "HK")),
            }
        )
    return {
        "path": path,
        "consultant": consultant,
        "client": client,
        "fiscal_year": fiscal_year,
        "primanota_number": primanota_number,
        "primanota_label": primanota_label,
        "period": period,
        "rows": data,
    }


def _fake_guid(pn: dict, nr: int) -> str:
    name = f"{pn['consultant']}/{pn['client']}/{pn['primanota_number']}/{nr}"
    return str(_uuid.uuid5(_GUID_NAMESPACE, name))


def _fake_batch_number(primanota_number: str) -> int:
    """Deterministic 6-digit stand-in for the undocumented header field 26."""
    digest = _hashlib.sha1(primanota_number.encode()).hexdigest()
    return int(digest, 16) % 900_000 + 100_000


def _write_fake_dtvf(pn: dict, out_path: _pathlib.Path) -> None:
    period = pn["period"]
    last_day = _calendar.monthrange(period.year, period.month)[1]
    datum_von = period
    datum_bis = period.replace(day=last_day)
    if pn["fiscal_year"] < 2026:
        sachkontenlaenge, sachkontenrahmen, branchenloesung = 4, "03", None
    else:
        sachkontenlaenge, sachkontenrahmen, branchenloesung = 5, "42", 4910
    hk_counts: dict[str, int] = {}
    for row in pn["rows"]:
        hk_counts[row["hk"] or "SV"] = hk_counts.get(row["hk"] or "SV", 0) + 1
    herkunft = max(hk_counts, key=lambda hk: hk_counts[hk]) if hk_counts else "SV"

    header_line = _datev.build_buchungsstapel_header_line(
        kennzeichen="DTVF",
        # Deterministic: end of the batch month, noon (naive local time, like
        # DATEV's own header timestamps). The real "erzeugt am" is not
        # reconstructible (and must not be relied on, see datev_import.md).
        erzeugt_am=_datetime.datetime(period.year, period.month, last_day, 12, 0, 0),  # noqa: DTZ001
        herkunft=herkunft,
        exportiert_von="wsjrdp-fake",
        beraternummer=pn["consultant"],
        mandantennummer=pn["client"],
        wj_beginn=_datetime.date(pn["fiscal_year"], 1, 1),
        sachkontenlaenge=sachkontenlaenge,
        datum_von=datum_von,
        datum_bis=datum_bis,
        bezeichnung=_cp1252_safe(pn["primanota_label"]) or pn["primanota_number"],
        derivatskennzeichen="MP",
        reserviert_26=_fake_batch_number(pn["primanota_number"]),
        sachkontenrahmen=sachkontenrahmen,
        branchenloesung_id=branchenloesung,
    )

    with open(out_path, "w", encoding="cp1252", newline="\r\n") as csvfile:
        # Serialization exactly like DATEV's own exporter (see DatevExtfWriter).
        writer = _datev.DatevExtfWriter(
            csvfile,
            type_overrides={"Beteiligtennummer": "Text", "Generalumkehr (GU)": "Zahl"},
        )
        csvfile.write(header_line + "\n")
        writer.writeheader()
        for row in pn["rows"]:
            record = {
                "Umsatz (ohne Soll/Haben-Kz)": _amount_de(row["umsatz"]),
                "Soll/Haben-Kennzeichen": row["sh"],
                "WKZ Umsatz": row["wkz"],
                "Kurs": row["kurs"],
                "Basis-Umsatz": row["umsatz_bw"],
                "WKZ Basis-Umsatz": row["wkz_bw"],
                "Konto": row["konto"],
                "Gegenkonto (ohne BU-Schlüssel)": row["gegenkonto"],
                "BU-Schlüssel": row["bu"],
                "Belegdatum": row["datum"].strftime("%d%m"),
                "Belegfeld 1": _cp1252_safe(row["belegfeld1"]),
                "Belegfeld 2": _cp1252_safe(row["belegfeld2"]),
                "Skonto": row["skonto"],
                "Buchungstext": _cp1252_safe(row["text"]),
                "KOST1 - Kostenstelle": row["kost1"],
                "KOST2 - Kostenstelle": row["kost2"],
                "Kost-Menge": row["kost_menge"],
                # Constant fills observed in every real DTVF row:
                "USt-Schlüssel (Anzahlungen)": 0,
                "Erlöskonto (Anzahlungen)": 0,
                "Herkunft-Kz": row["hk"] or "SV",
                "Buchungs GUID": _fake_guid(pn, row["nr"]),
                "Skontosperre": 0,
                "Festschreibung": 0,
                "Generalumkehr (GU)": 1 if row["gu"] else 0,
                "BVV-Position": 0,
            }
            if row["leistungsdatum"] not in (None, ""):
                record["Leistungsdatum"] = _parse_date(row["leistungsdatum"]).strftime(
                    "%d%m%Y"
                )
            writer.writerow(record)


def main(argv: list[str] | None = None) -> int:
    _logging.basicConfig(level=_logging.INFO, format="%(levelname)s %(message)s")
    parser = _argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--glob",
        default="External_Data/DATEV_Primanota_Examples/118531_66352_2025_*.xlsx",
        help="Primanota Excel files to convert (default: the 2025 sheets)",
    )
    parser.add_argument(
        "--out-dir",
        default="External_Data/Fake_DATEV_Export_2025",
        help="Output directory (default: External_Data/Fake_DATEV_Export_2025)",
    )
    parser.add_argument(
        "--stamp",
        default=None,
        help="Export stamp for the file names (YYYYMMDD_HHMMSS); default: "
        "<max fiscal year>1231_000000 (deterministic)",
    )
    args = parser.parse_args(argv)

    paths = sorted(_pathlib.Path(".").glob(args.glob))
    if not paths:
        parser.error(f"No files match {args.glob!r}")
    primanotas = [_read_primanota(p) for p in paths]
    primanotas.sort(
        key=lambda pn: (pn["fiscal_year"], pn["period"], pn["primanota_number"])
    )

    stamp = args.stamp or f"{max(pn['fiscal_year'] for pn in primanotas)}1231_000000"
    out_dir = _pathlib.Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    for seq, pn in enumerate(primanotas, start=1):
        out_path = out_dir / f"DTVF_Buchungsstapel_{stamp}_{seq:05d}.csv"
        _write_fake_dtvf(pn, out_path)
        _LOGGER.info(
            "%s  <-  %s (%s %r, %d Zeilen)",
            out_path,
            pn["path"].name,
            pn["primanota_number"],
            pn["primanota_label"],
            len(pn["rows"]),
        )
    _LOGGER.info("Fertig: %d Stapel nach %s", len(primanotas), out_dir)
    return 0


if __name__ == "__main__":
    import sys

    sys.exit(main())
