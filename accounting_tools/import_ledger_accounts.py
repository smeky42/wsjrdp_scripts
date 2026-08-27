#!/usr/bin/env -S uv run
"""Import ledger-account (Sachkonto) master data into wsjrdp_ledger_accounts,
auto-detecting the source of each given file.

Supported sources (detected by extension + header):

  * DATEV Sachkonten export (.xlsx, or .csv with a "Konto von" column) --
    writes: name, account_type (derived from the number for now), datev_purpose
    ("Kontenzweck"), datev_function_type ("HFTyp"), datev_function_number
    ("Funktion"), datev_additional_function ("Zusatzfunktion") and any remaining
    populated columns into other_datev_columns (FE, Anlagenspiegelfkt.,
    non-standard S/B/K/I herkunft, ...).
  * Moss subcategories export (.csv) -- columns
    "Name","Expense Account - Number","Category","Status". Writes: name,
    moss_status, moss_category, and any extra columns into other_moss_columns.
  * Hitobito accounts CSV -- columns number,name,short_name. Writes: name,
    short_name (Hitobito-only), account_type (derived from the number).

Rules:
  * `number` is the shared key; account numbers are left-padded to 5 digits
    (Excel/Moss drop the leading zero of class-0 accounts).
  * Every source MAY overwrite `name`. Moss and DATEV names are expected to be
    identical (kept in sync); a mismatch against the stored name is logged as a
    likely sync error, then overwritten.
  * Each branch touches ONLY its own columns. The DATEV branch never changes
    moss_*/other_moss_columns; the Moss branch never changes
    datev_*/account_type/other_datev_columns; short_name is written only by the
    Hitobito branch; no branch touches the Hitobito-only aliases,
    description, comment, visibility and additional_info (display_short_name
    is database-generated; moss_status stays NULL for accounts Moss does not
    know).
  * Before writing, the stored rows are loaded (read-only) and diffed against
    the incoming data (CP1252-transliteration-aware for DATEV text). Accounts
    whose target state equals the stored state are NOT touched at all; the
    remainder is split into explicit INSERTs (unknown number) and UPDATEs
    (no upsert). Idempotent: a re-run of the same files touches nothing.
  * The other_*_columns JSONB is MERGED per key: a file is authoritative only
    for the keys its format covers -- covered keys are set (or DELETED when
    the export carries no value for them), keys from other/older sources are
    preserved.
  * 6-digit numbers (Debitoren 1xxxxx-6xxxxx, Kreditoren 7xxxxx-9xxxxx) are a
    hard error in every source (they are personal accounts ->
    wsjrdp_personal_accounts; a CHECK constraint rejects them).
"""

from __future__ import annotations

import csv as _csv
import logging as _logging
import pathlib as _pathlib
import re as _re
import sys as _sys
import typing as _typing

import wsjrdp2027
from wsjrdp2027 import SpecialValue
from wsjrdp2027._internal.single_table_upsert_plan import (
    SingleTableUpsertPlanBuilder,
)


_SELF_NAME = _pathlib.Path(__file__).stem
_LOGGER = _logging.getLogger(__name__)

_TABLE_NAME = "wsjrdp_ledger_accounts"

# Column order of the DATEV Sachkonten export (positional -- the three "S/B/K/I"
# herkunft columns share a header, so we cannot read it as a dict).
_DATEV_HEADER = [
    "Konto von",
    "Konto bis",
    "S/B/K/I",
    "Beschriftung",
    "S/B/K/I",
    "Zusatzfunktion",
    "HFTyp",
    "Funktion",
    "FE",
    "Faktor 2",
    "Konto 1",
    "Konto 2",
    "S/B/K/I",
    "Anlagenspiegelfkt.",
    "S/B/I",
    "Kontenzweck",
]


def _norm(value: object) -> str:
    return "" if value is None else str(value).strip()


def _account_number(value: object) -> str:
    """Trimmed account number, left-padded to 5 digits (Excel/Moss drop the
    leading zero of class-0 accounts). Non-numeric stays as-is."""
    s = _norm(value)
    return s.zfill(5) if s.isdigit() and len(s) < 5 else s


def _int_or_none(value: object) -> int | None:
    s = _norm(value)
    return int(s) if s.lstrip("-").isdigit() else None


def _check_no_personal_accounts(numbers: _typing.Iterable[str], source: str) -> None:
    """Six-digit numbers are personal accounts (Debitoren 1xxxxx-6xxxxx,
    Kreditoren 7xxxxx-9xxxxx); the chk_ledger_account_number_not_personal_account
    CHECK constraint rejects them. Abort up front rather than failing
    mid-transaction."""
    personal = sorted({n for n in numbers if _re.fullmatch(r"[1-9]\d{5}", n)})
    if personal:
        _LOGGER.error(
            "%s: %d personal account number(s) (six-digit): %s",
            source,
            len(personal),
            ", ".join(personal[:10]),
        )
        raise SystemExit(
            f"Import aborted ({source}): six-digit numbers (Debitoren/Kreditoren) "
            "belong to wsjrdp_personal_accounts (import_personal_accounts.py), not "
            "ledger accounts. Re-export the Sachkonten WITHOUT bespielte Personenkonten."
        )


# --- source detection --------------------------------------------------------


def _iter_rows(path: _pathlib.Path) -> _typing.Iterator[list]:
    """Yield each row as a positional list, for .xlsx or .csv alike."""
    if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for row in wb[wb.sheetnames[0]].iter_rows(values_only=True):
                yield list(row)
        finally:
            wb.close()
    else:
        with path.open(encoding="utf-8-sig", newline="") as f:
            for row in _csv.reader(f):
                yield row


def _detect_source(path: _pathlib.Path) -> str:
    if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        return "datev"
    with path.open(encoding="utf-8-sig", newline="") as f:
        header = f.readline()
    if "Expense Account - Number" in header:
        return "moss"
    if "Konto von" in header:
        return "datev"
    if "short_name" in header:
        return "hitobito"
    raise SystemExit(f"{path}: unbekanntes Stammdaten-Format (Header: {header!r})")


# --- readers (return list of column dicts) -----------------------------------


def _read_datev(path: _pathlib.Path) -> list[dict]:
    rows = _iter_rows(path)
    header = [_norm(c) for c in next(rows, [])]
    if len(header) < 16 or header[0] != "Konto von" or header[15] != "Kontenzweck":
        raise SystemExit(
            f"{path}: unerwartete DATEV-Sachkonten-Spalten (erwartet 'Konto von' "
            f"... 'Kontenzweck', gefunden: {header[:2]} ... {header[-1:]})"
        )
    out: list[dict] = []
    for r in rows:
        r = list(r) + [None] * (16 - len(r))
        number = _account_number(r[0])
        if not number:
            continue
        konto_bis = _account_number(r[1])
        if konto_bis and konto_bis != number:
            _LOGGER.warning(
                "Konto %s: Bereich Konto von != Konto bis (%s..%s) -- als "
                "Einzelkonto %s behandelt",
                number,
                number,
                konto_bis,
                number,
            )
        # The export is authoritative for exactly these keys: a meaningful
        # value is set, an empty/default one DELETES the key; keys from other
        # sources inside the JSONB stay untouched (SpecialValue.DELETE semantics).
        other: dict[str, object] = {}
        for idx, key in (
            (2, "S/B/K/I (Beschriftung)"),
            (4, "S/B/K/I (Kontenfunktion)"),
            (12, "S/B/K/I (Anlagenspiegel)"),
            (14, "S/B/I (Kontenzweck)"),
        ):
            v = _norm(r[idx])
            other[key] = v if (v and v != "S") else SpecialValue.DELETE
        for idx, key in ((8, "FE"), (13, "Anlagenspiegelfkt.")):
            n = _int_or_none(r[idx])
            other[key] = n if n else SpecialValue.DELETE
        out.append(
            {
                "number": number,
                "name": _norm(r[3]) or None,
                "account_type": wsjrdp2027.datev.account_type_for_account_number(
                    number
                ),
                "datev_purpose": _norm(r[15]) or None,
                "datev_function_type": _int_or_none(r[6]),  # HFTyp
                "datev_function_number": _int_or_none(r[7]),  # Funktion
                "datev_additional_function": _int_or_none(r[5]),  # Zusatzfunktion
                "other_datev_columns": other,
            }
        )
    return out


_MOSS_KNOWN = {"Name", "Expense Account - Number", "Category", "Status"}


def _read_moss(path: _pathlib.Path) -> list[dict]:
    out: list[dict] = []
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = _csv.DictReader(f)
        extra_headers = [
            h for h in (reader.fieldnames or []) if h and h not in _MOSS_KNOWN
        ]
        for row in reader:
            number = _account_number(row.get("Expense Account - Number"))
            if not number:
                continue
            status = _norm(row.get("Status")).lower() or "active"
            # The export is authoritative for its extra headers: empty value
            # -> the key is deleted from the stored JSONB (SpecialValue.DELETE).
            other = {
                h: (_norm(row.get(h)) or SpecialValue.DELETE) for h in extra_headers
            }
            out.append(
                {
                    "number": number,
                    "name": _norm(row.get("Name")) or None,
                    "moss_status": "deactivated"
                    if status == "deactivated"
                    else "active",
                    "moss_category": _norm(row.get("Category")) or None,
                    "other_moss_columns": other,
                }
            )
    return out


def _read_hitobito(path: _pathlib.Path) -> list[dict]:
    out: list[dict] = []
    with path.open(encoding="utf-8-sig", newline="") as f:
        for row in _csv.DictReader(f):
            number = _account_number(row.get("number"))
            if not number:
                continue
            out.append(
                {
                    "number": number,
                    "name": _norm(row.get("name")) or None,
                    "short_name": _norm(row.get("short_name")) or None,
                    "account_type": wsjrdp2027.datev.account_type_for_account_number(
                        number
                    ),
                }
            )
    return out


# --- diff / plan / write -----------------------------------------------------


# Free-text columns the DATEV branch writes. A DATEV export carries only the
# CP1252-representable form of their content (verified: the Sachkonten export
# contains no character outside CP1252), so they are compared through
# to_win1252_compatible instead of directly -- otherwise a DATEV import would
# overwrite a richer Unicode value coming from Moss. See doc/fin/datev_cp1252.md
# in the hitobito_wsjrdp_2027 wagon.
_CP1252_PROTECTED = ("name", "datev_purpose")


def _warn_name_drift(updates: list[dict], existing: dict) -> None:
    """A name inside an update delta is a genuine change (transliteration-equal
    values never reach the delta): where a name was stored before, the sources
    drifted apart."""
    drift = [
        (u["number"], existing.get(u["number"], {}).get("name"), u["name"])
        for u in updates
        if u.get("name") and existing.get(u["number"], {}).get("name") is not None
    ]
    for number, old, new in drift[:20]:
        _LOGGER.warning(
            "Namens-Drift Konto %s: gespeichert %r != importiert %r (Sync-Fehler?)",
            number,
            old,
            new,
        )
    if drift:
        _LOGGER.warning(
            "%d Konto/-en mit abweichendem Namen (Name wird ueberschrieben).",
            len(drift),
        )


def _log_plan_summary(planned) -> None:
    """Show what an approval would apply -- logged BEFORE the production
    approval is requested."""
    for table, counts in planned.operation_counts().items():
        _LOGGER.info(
            "Geplant fuer %s: %d INSERTs, %d UPDATEs, %d DELETEs (%d unangetastet).",
            table,
            counts.inserts,
            counts.updates,
            counts.deletes,
            len(planned.untouched_keys),
        )
    for label, rows in (("INSERT", planned.inserts), ("UPDATE", planned.updates)):
        if not rows:
            continue
        keys = [str(row["number"]) for row in rows]
        shown = ", ".join(keys[:10])
        if len(keys) > 10:
            shown += f", ... (+{len(keys) - 10} weitere)"
        _LOGGER.info("  %s: %s", label, shown)


# reader + columns compared CP1252-transliteration-aware (DATEV text fields).
_BRANCHES: dict[str, tuple] = {
    "datev": (_read_datev, ("name", "datev_purpose", "other_datev_columns")),
    "moss": (_read_moss, False),
    "hitobito": (_read_hitobito, False),
}


def create_argument_parser():
    import argparse

    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument(
        "files",
        nargs="+",
        help="Stammdaten-Dateien (DATEV .xlsx/.csv, Moss .csv, Hitobito .csv); "
        "die Quelle wird je Datei automatisch erkannt.",
    )
    p.add_argument(
        "--rollback-for-testing",
        action="store_true",
        default=False,
        help="Apply the plan, then ROLLBACK instead of committing (testing).",
    )
    return p


def main(argv=None):
    ctx = wsjrdp2027.WsjRdpContext(
        argument_parser=create_argument_parser(), argv=argv, __file__=__file__
    )
    out_base = ctx.make_out_path(_SELF_NAME + "_{{ filename_suffix }}")
    ctx.configure_log_file(out_base.with_suffix(".log"))

    # All files are merged into ONE plan builder (in argument order; a later
    # file wins on colliding columns, and for DATEV files a mere CP1252
    # transliteration never overwrites an earlier Unicode value), so there is
    # exactly one plan and at most one apply() at the end.
    builder = SingleTableUpsertPlanBuilder(
        _TABLE_NAME, "number", [], time_zone=ctx.hitobito_time_zone
    )
    plan_translit: set[str] = set()
    for path_str in ctx.parsed_args.files:
        path = _pathlib.Path(path_str)
        source = _detect_source(path)
        reader, translit_columns = _BRANCHES[source]
        rows = reader(path)
        _check_no_personal_accounts(
            (r["number"] for r in rows), f"{source}:{path.name}"
        )
        _LOGGER.info(
            "%s: %d Konten aus %s (Quelle: %s)", path.name, len(rows), path, source
        )
        if translit_columns:
            file_columns = {column for row in rows for column in row}
            columns = [c for c in translit_columns if c in file_columns]
            plan_translit.update(columns)
            builder.merge_values(rows, keep_existing_for_cp1252_equality=columns)
        else:
            builder.merge_values(rows)

    with ctx:
        ro_conn = ctx.hitobito_psycopg_connection(read_only=True)
        builder.load_existing(ro_conn)
        planned = builder.plan(
            skip_update_for_cp1252_equality=sorted(plan_translit) or False
        )
        _warn_name_drift(planned.updates, builder.existing)

        _log_plan_summary(planned)

        if not planned.inserts and not planned.updates:
            _LOGGER.info(
                "Sachkonten: nichts zu schreiben (%d unangetastet).",
                len(planned.untouched_keys),
            )
            return

        if ctx.dry_run:
            _LOGGER.info("[dry-run] Plan wird nicht angewendet.")
            return

        # The plan summary above shows exactly what this approval applies.
        ctx.require_approval_to_run_in_prod()
        rw_conn = ctx.hitobito_psycopg_connection(read_only=False)
        # Aware start time; the pg_table_* helpers turn it into
        # Rails-compatible UTC-naive created_at/updated_at stamps via the
        # session time zone.
        inserted, updated = planned.apply(rw_conn, now=ctx.start_time)
        _LOGGER.info(
            "Sachkonten: %d neu, %d aktualisiert, %d unangetastet (identisch).",
            len(inserted),
            len(updated),
            len(planned.untouched_keys),
        )
        if ctx.parsed_args.rollback_for_testing:
            _LOGGER.warning(
                "ROLLBACK (--rollback-for-testing given) - no changes committed"
            )
            rw_conn.rollback()
        # The commit happens implicitly when the `with ctx:` block exits
        # cleanly; an exception before that leaves the database untouched.


if __name__ == "__main__":
    _sys.exit(main())
